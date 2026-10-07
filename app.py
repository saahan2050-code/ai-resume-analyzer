"""ATS Resume Checker - Streamlit + Gemini Flash.

Upload a resume (PDF or DOCX), optionally paste a job description, and get an
ATS score with concrete improvements.
"""

import io
import json
import os
import re
import time

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
DEFAULT_MODEL = "gemini-3.5-flash"
FALLBACK_MODELS = ["gemini-3.5-flash-lite"]  # tried if the main model stays overloaded
RETRIES_PER_MODEL = 3  # attempts per model, with exponential backoff
RETRYABLE_CODES = {429, 500, 502, 503, 504}
MAX_FILE_MB = 5
MAX_CHARS = 30_000  # keeps the prompt small and fast
MIN_CHARS = 150  # below this we assume the PDF is a scanned image

SYSTEM_PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and \
professional resume reviewer. Evaluate the resume text provided and respond \
with ONLY a valid JSON object (no markdown fences, no commentary) matching \
this exact structure:

{
  "overall_score": <integer 0-100>,
  "summary": "<2-3 sentence overall assessment>",
  "category_scores": {
    "keywords": <integer 0-100>,
    "formatting": <integer 0-100>,
    "experience_impact": <integer 0-100>,
    "skills": <integer 0-100>,
    "education": <integer 0-100>,
    "readability": <integer 0-100>
  },
  "strengths": ["<string>", "..."],
  "missing_keywords": ["<string>", "..."],
  "formatting_issues": ["<string>", "..."],
  "improvements": [
    {
      "priority": "high" | "medium" | "low",
      "section": "<resume section, e.g. Summary, Experience, Skills>",
      "issue": "<what is wrong>",
      "suggestion": "<specific fix>",
      "example": "<rewritten example line, or empty string>"
    }
  ]
}

Scoring guidance:
- Judge ATS-friendliness: standard section headings, parseable structure, \
relevant keywords, quantified achievements, action verbs, consistent dates, \
contact info, appropriate length, and no spelling or grammar problems.
- If a job description is provided, weigh keyword and skills match against it \
heavily and list the important missing keywords from it.
- If no job description is provided, judge against general best practices for \
the role the resume appears to target.
- Be honest and strict. Do not inflate scores. Most resumes score 50-80.
- Give 5-10 improvements, ordered from highest to lowest priority.
- Only reference content that actually appears in the resume. Never invent \
experience, employers, or numbers; in "example", use placeholders like [X%] \
where a number is needed.
"""


# ----------------------------------------------------------------------------
# Text extraction
# ----------------------------------------------------------------------------
def extract_text_from_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ValueError("This PDF is password protected.")
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages)


def extract_text_from_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # Many resumes keep content inside tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text.strip())
    return "\n".join(parts)


def extract_resume_text(filename: str, data: bytes) -> str:
    name = filename.lower()
    if name.endswith(".pdf"):
        text = extract_text_from_pdf(data)
    elif name.endswith(".docx"):
        text = extract_text_from_docx(data)
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# ----------------------------------------------------------------------------
# Gemini
# ----------------------------------------------------------------------------
def get_api_key(sidebar_key: str) -> str:
    if sidebar_key.strip():
        return sidebar_key.strip()
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass  # no secrets file locally
    return os.environ.get("GEMINI_API_KEY", "")


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating stray markdown fences."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        start, end = raw.find("{"), raw.rfind("}")
        if start != -1 and end > start:
            return json.loads(raw[start : end + 1])
        raise


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def _str_list(value) -> list:
    if not isinstance(value, list):
        return []
    return [str(v).strip() for v in value if str(v).strip()]


def normalize_result(data: dict) -> dict:
    """Make sure the model output has every field the UI needs."""
    if not isinstance(data, dict):
        raise ValueError("Model returned an unexpected format.")
    cats = data.get("category_scores") or {}
    improvements = []
    for item in data.get("improvements") or []:
        if not isinstance(item, dict):
            continue
        priority = str(item.get("priority", "medium")).lower()
        if priority not in ("high", "medium", "low"):
            priority = "medium"
        improvements.append(
            {
                "priority": priority,
                "section": str(item.get("section", "General")),
                "issue": str(item.get("issue", "")),
                "suggestion": str(item.get("suggestion", "")),
                "example": str(item.get("example", "") or ""),
            }
        )
    order = {"high": 0, "medium": 1, "low": 2}
    improvements.sort(key=lambda x: order[x["priority"]])
    return {
        "overall_score": _clamp(data.get("overall_score")),
        "summary": str(data.get("summary", "")),
        "category_scores": {
            "Keywords": _clamp(cats.get("keywords")),
            "Formatting": _clamp(cats.get("formatting")),
            "Experience impact": _clamp(cats.get("experience_impact")),
            "Skills": _clamp(cats.get("skills")),
            "Education": _clamp(cats.get("education")),
            "Readability": _clamp(cats.get("readability")),
        },
        "strengths": _str_list(data.get("strengths")),
        "missing_keywords": _str_list(data.get("missing_keywords")),
        "formatting_issues": _str_list(data.get("formatting_issues")),
        "improvements": improvements,
    }


def _is_retryable(err: Exception) -> bool:
    code = getattr(err, "code", None)
    if isinstance(code, int):
        return code in RETRYABLE_CODES
    msg = str(err)
    return any(tok in msg for tok in ("503", "UNAVAILABLE", "429", "RESOURCE_EXHAUSTED", "overloaded"))


def _generate_with_retry(client, models, prompt, config, sleep=time.sleep):
    """Try each model in order; retry transient errors with exponential backoff."""
    last_err = None
    for model in models:
        for attempt in range(RETRIES_PER_MODEL):
            try:
                return client.models.generate_content(
                    model=model, contents=prompt, config=config
                )
            except Exception as e:  # noqa: BLE001
                last_err = e
                if not _is_retryable(e):
                    raise
                if attempt < RETRIES_PER_MODEL - 1:
                    sleep(2 ** attempt * 2)  # 2s, 4s
        # all attempts for this model failed -> fall through to next model
    raise last_err


def analyze_resume(api_key: str, model: str, resume_text: str, job_desc: str) -> dict:
    client = genai.Client(api_key=api_key)
    prompt = f"RESUME:\n\"\"\"\n{resume_text[:MAX_CHARS]}\n\"\"\"\n"
    if job_desc.strip():
        prompt += f"\nJOB DESCRIPTION:\n\"\"\"\n{job_desc.strip()[:MAX_CHARS]}\n\"\"\"\n"
    else:
        prompt += "\nJOB DESCRIPTION: (not provided)\n"

    config = types.GenerateContentConfig(
        system_instruction=SYSTEM_PROMPT,
        response_mime_type="application/json",
        temperature=0.2,
    )
    models = [model] + [m for m in FALLBACK_MODELS if m != model]
    response = _generate_with_retry(client, models, prompt, config)
    return normalize_result(parse_json_response(response.text))


@st.cache_data(show_spinner=False, ttl=3600)
def cached_analysis(api_key: str, model: str, resume_text: str, job_desc: str) -> dict:
    # Same resume + JD + model returns the same result without a new API call
    return analyze_resume(api_key, model, resume_text, job_desc)


# ----------------------------------------------------------------------------
# UI helpers
# ----------------------------------------------------------------------------
def score_label(score: int):
    if score >= 80:
        return "Excellent", "green"
    if score >= 65:
        return "Good", "orange"
    if score >= 50:
        return "Needs work", "orange"
    return "Poor", "red"


PRIORITY_ICON = {"high": "🔴 High", "medium": "🟠 Medium", "low": "🟢 Low"}


def render_results(result: dict) -> None:
    score = result["overall_score"]
    label, color = score_label(score)

    col1, col2 = st.columns([1, 2])
    with col1:
        st.metric("ATS Score", f"{score} / 100")
        st.markdown(f":{color}[**{label}**]")
    with col2:
        st.progress(score / 100)
        if result["summary"]:
            st.write(result["summary"])

    st.subheader("Score breakdown")
    cols = st.columns(3)
    for i, (name, val) in enumerate(result["category_scores"].items()):
        with cols[i % 3]:
            st.write(f"**{name}** - {val}/100")
            st.progress(val / 100)

    left, right = st.columns(2)
    with left:
        st.subheader("✅ Strengths")
        for s in result["strengths"] or ["No specific strengths returned."]:
            st.markdown(f"- {s}")
    with right:
        st.subheader("🔑 Missing keywords")
        if result["missing_keywords"]:
            st.markdown(" ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.write("None identified.")

    if result["formatting_issues"]:
        st.subheader("⚠️ Formatting issues")
        for f in result["formatting_issues"]:
            st.markdown(f"- {f}")

    st.subheader("🛠️ Recommended improvements")
    for imp in result["improvements"]:
        title = f"{PRIORITY_ICON[imp['priority']]} · {imp['section']}"
        with st.expander(title, expanded=imp["priority"] == "high"):
            st.markdown(f"**Issue:** {imp['issue']}")
            st.markdown(f"**Fix:** {imp['suggestion']}")
            if imp["example"]:
                st.markdown("**Example:**")
                st.code(imp["example"], language=None)

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


# ----------------------------------------------------------------------------
# App
# ----------------------------------------------------------------------------
def main() -> None:
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="wide")
    st.title("📄 ATS Resume Checker")
    st.caption("Upload your resume to get an ATS score and specific ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        sidebar_key = st.text_input(
            "Gemini API key",
            type="password",
            help="Optional if GEMINI_API_KEY is set in secrets or environment.",
        )
        model = st.text_input("Model", value=DEFAULT_MODEL)
        st.markdown("[Get a free API key](https://aistudio.google.com/apikey)")
        st.info("Your resume is sent to Google's Gemini API for analysis.")

    uploaded = st.file_uploader("Upload resume (PDF or DOCX)", type=["pdf", "docx"])
    job_desc = st.text_area(
        "Job description (optional, improves keyword matching)",
        height=150,
        placeholder="Paste the job description here...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        api_key = get_api_key(sidebar_key)
        if not api_key:
            st.error("Please provide a Gemini API key in the sidebar.")
            st.stop()

        data = uploaded.getvalue()
        if len(data) > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            st.stop()

        try:
            with st.spinner("Reading resume..."):
                text = extract_resume_text(uploaded.name, data)
        except Exception as e:
            st.error(f"Could not read the file: {e}")
            st.stop()

        if len(text) < MIN_CHARS:
            st.error(
                "Very little text could be extracted. If your resume is a scanned "
                "image, ATS systems cannot read it either - export it as a "
                "text-based PDF or DOCX and try again."
            )
            st.stop()

        try:
            with st.spinner("Analyzing with Gemini..."):
                result = cached_analysis(api_key, model.strip() or DEFAULT_MODEL, text, job_desc)
        except json.JSONDecodeError:
            st.error("The AI returned an unreadable response. Please try again.")
            st.stop()
        except Exception as e:
            if _is_retryable(e):
                st.error(
                    "Google's Gemini service is overloaded right now (tried several "
                    "times and a backup model). Please wait a minute and click "
                    "Analyze again."
                )
            else:
                st.error(f"Analysis failed: {e}")
            st.stop()

        st.session_state["result"] = result

    if "result" in st.session_state:
        st.divider()
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
